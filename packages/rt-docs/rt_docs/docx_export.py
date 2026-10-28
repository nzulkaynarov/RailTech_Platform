"""Тело JSON → DOCX («Скачать DOCX» для юриста контрагента). DOCX — формат обмена, не формат шаблона (ADR-0020 п. 3).

Библиотека — python-docx (MIT, чистый Python поверх lxml): без Office и LibreOffice на сервере (образ и CPU qemu64 не
меняются), документ собирается из того же разрешённого дерева, что и PDF (rt_docs.resolve). Что не переносится 1:1:
разбивка на страницы (её делает Word), правило «Стр. N из M только при M > 1» (поле Word ставится, если вызывающий
знает, что страниц больше одной — `page_numbers`), «Исполнитель» внизу последней страницы (печатается последним
абзацем), логотип WebP (python-docx его не вставляет — предупреждение).
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .contracts import Catalog, load_catalog, validate_body
from .css import PageSetup, page_setup
from .errors import ContractsNotFoundError
from .formatters import format_date
from .labels import EXECUTOR, MISSING_BLOCK, PAGE_OF, REPLY_TO, SIGN, STAMP, number_line
from .model import Density, DocumentMeta, Lang, Layout, Letterhead, Orientation, Tuning, Value
from .resolve import resolve

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_FONT = {"sans": "Noto Sans", "serif": "Noto Serif"}
_MUTED = RGBColor(0x5A, 0x62, 0x70)


@dataclass(frozen=True)
class DocxResult:
    data: bytes
    warnings: tuple[str, ...] = ()


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


class _DocxWriter:
    def __init__(self, document: DocxDocument, lang: Lang, letterhead: Letterhead, setup: PageSetup) -> None:
        self.doc = document
        self.lang = lang
        self.letterhead = letterhead
        self.setup = setup
        self.accent = _rgb(letterhead.accent)

    # ─── inline ──────────────────────────────────────────────────────────────────────────────────

    def inline(self, paragraph: Paragraph, nodes: Sequence[Mapping[str, Any]]) -> None:
        for node in nodes:
            if node["type"] == "hardBreak":
                paragraph.add_run().add_break()
                continue
            marks = {m["type"]: m for m in node.get("marks") or ()}
            text = f"[{node['text']}]" if node.get("_missing") else str(node["text"])
            if "link" in marks:
                self._hyperlink(paragraph, text, str(marks["link"]["attrs"]["href"]), marks)
                continue
            run = paragraph.add_run(text)
            run.bold = "bold" in marks or None
            run.italic = "italic" in marks or None
            run.underline = "underline" in marks or None
            if node.get("_missing"):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def _hyperlink(self, paragraph: Paragraph, text: str, href: str, marks: Mapping[str, Any]) -> None:
        rel = paragraph.part.relate_to(
            href, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        props.append(underline)
        if "bold" in marks:
            props.append(OxmlElement("w:b"))
        if "italic" in marks:
            props.append(OxmlElement("w:i"))
        run.append(props)
        text_el = OxmlElement("w:t")
        text_el.text = text
        text_el.set(qn("xml:space"), "preserve")
        run.append(text_el)
        link.append(run)
        paragraph._p.append(link)

    # ─── блоки ───────────────────────────────────────────────────────────────────────────────────

    def blocks(self, nodes: Sequence[Mapping[str, Any]], container: Any = None, level: int = 0) -> None:
        for node in nodes:
            self.block(node, container if container is not None else self.doc, level)

    def block(self, node: Mapping[str, Any], container: Any, level: int) -> None:
        kind = node["type"]
        attrs = node.get("attrs") or {}
        if kind == "paragraph":
            paragraph = container.add_paragraph()
            if attrs.get("textAlign") in _ALIGN:
                paragraph.alignment = _ALIGN[attrs["textAlign"]]
            self.inline(paragraph, node.get("content", ()))
        elif kind == "heading":
            paragraph = container.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(self.setup.font_pt * 0.9)
            if attrs.get("textAlign") in _ALIGN:
                paragraph.alignment = _ALIGN[attrs["textAlign"]]
            self.inline(paragraph, node.get("content", ()))
            scale = {1: 1.4, 2: 1.18, 3: 1.04}[int(attrs["level"])]
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(round(self.setup.font_pt * scale, 1))
        elif kind in ("bulletList", "orderedList"):
            base = "List Bullet" if kind == "bulletList" else "List Number"
            style = base if level == 0 else f"{base} {min(level + 1, 3)}"
            for item in node["content"]:
                bullet_done = False
                for child in item["content"]:
                    if child["type"] == "paragraph" and not bullet_done:
                        paragraph = container.add_paragraph(style=style)
                        self.inline(paragraph, child.get("content", ()))
                        bullet_done = True
                    else:
                        self.block(child, container, level + 1 if child["type"] != "paragraph" else level)
        elif kind == "table":
            self.table(node, container)
        elif kind == "pageBreak":
            container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "signature":
            self.signature(node, container)
        elif kind == "_block":
            if node.get("missing"):
                paragraph = container.add_paragraph()
                run = paragraph.add_run(MISSING_BLOCK[self.lang].format(label=node["label"]))
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            else:
                self.blocks(node["content"], container, level)

    def table(self, node: Mapping[str, Any], container: Any) -> None:
        rows: list[Mapping[str, Any]] = list(node["content"])
        attrs = node.get("attrs") or {}
        widths = attrs.get("columnWidths")
        columns = max(sum(int((c.get("attrs") or {}).get("colspan", 1)) for c in r["content"]) for r in rows)
        table: Table = container.add_table(rows=len(rows), cols=columns)
        variant = attrs.get("variant", "grid")
        if variant == "grid":
            table.style = "Table Grid"
        text_width = self._text_width_mm()
        if widths and len(widths) == columns:
            total = sum(float(w) for w in widths) or 1.0
            for index, column in enumerate(table.columns):
                column.width = Mm(text_width * float(widths[index]) / total)
        occupied: set[tuple[int, int]] = set()
        header_done = False
        for r_index, row_node in enumerate(rows):
            is_header = all(c["type"] == "tableHeader" for c in row_node["content"])
            if is_header and not header_done:
                _repeat_header(table.rows[r_index])
            else:
                header_done = True
            _cant_split(table.rows[r_index])
            c_index = 0
            for cell_node in row_node["content"]:
                while (r_index, c_index) in occupied:
                    c_index += 1
                if c_index >= columns:
                    break
                cattrs = cell_node.get("attrs") or {}
                colspan = int(cattrs.get("colspan", 1))
                rowspan = int(cattrs.get("rowspan", 1))
                cell: _Cell = table.cell(r_index, c_index)
                end_r = min(r_index + rowspan, len(rows)) - 1
                end_c = min(c_index + colspan, columns) - 1
                if end_r > r_index or end_c > c_index:
                    cell = cell.merge(table.cell(end_r, end_c))
                for rr in range(r_index, end_r + 1):
                    for cc in range(c_index, end_c + 1):
                        occupied.add((rr, cc))
                self._cell(cell, cell_node, cattrs, emphasis=bool((row_node.get("attrs") or {}).get("emphasis")))
                c_index = end_c + 1

    def _cell(self, cell: _Cell, node: Mapping[str, Any], attrs: Mapping[str, Any], *, emphasis: bool) -> None:
        cell.text = ""
        first = True
        for child in node["content"]:
            paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            align = "right" if attrs.get("numeric") else attrs.get("align")
            if align in _ALIGN:
                paragraph.alignment = _ALIGN[align]
            self.inline(paragraph, child.get("content", ()))
            for run in paragraph.runs:
                if node["type"] == "tableHeader" or emphasis:
                    run.bold = True

    def signature(self, node: Mapping[str, Any], container: Any) -> None:
        parties: list[Mapping[str, Any]] = list(node["parties"])
        table = container.add_table(rows=1, cols=2)
        for index, party in enumerate(parties):
            cell = table.cell(0, index)
            cell.text = ""
            lines: list[tuple[str, bool]] = []
            if party.get("role"):
                lines.append((self._field(party["role"]), True))
            if party.get("position"):
                lines.append((self._field(party["position"]), False))
            lines.append(("_" * 28, False))
            name = self._field(party["name"]) if party.get("name") else ""
            lines.append((SIGN[self.lang] + (f" · {name}" if name else ""), False))
            if node.get("stamp"):
                lines.append((STAMP[self.lang], False))
            for i, (value, bold) in enumerate(lines):
                paragraph = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                paragraph.paragraph_format.keep_with_next = True
                run = paragraph.add_run(value)
                run.bold = bold or None

    def _field(self, value: tuple[str, bool]) -> str:
        text, missing = value
        return f"[{text}]" if missing else text

    def _text_width_mm(self) -> float:
        _, right, _, left = self.setup.margins_mm
        width = 297 if self.setup.orientation is Orientation.LANDSCAPE else 210
        return float(width - left - right)


def _repeat_header(row: Any) -> None:
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    props.append(element)


def _cant_split(row: Any) -> None:
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    props.append(element)


def _field_run(paragraph: Paragraph, instruction: str) -> None:
    """Поле Word (PAGE, NUMPAGES): begin → instrText → separate → «1» → end."""
    for kind, text in (("begin", None), ("instr", instruction), ("separate", None), ("text", "1"), ("end", None)):
        run = paragraph.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
            run._r.append(el)
        elif kind == "text":
            run.text = str(text)
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


def to_docx(
    body: Mapping[str, Any],
    *,
    letterhead: Letterhead,
    layout: Layout,
    meta: DocumentMeta | None = None,
    values: Mapping[str, Value] | None = None,
    blocks: Mapping[str, Sequence[Mapping[str, Any]]] | None = None,
    catalog: Catalog | None = None,
    page_numbers: bool = True,
    contracts_dir: str | None = None,
) -> DocxResult:
    """DOCX из того же тела и бланка, что PDF.

    `page_numbers=False` — когда известно, что страница одна (статистика PDF).
    """
    validate_body(body, contracts_dir)
    if catalog is None:
        try:
            catalog = load_catalog(contracts_dir)
        except ContractsNotFoundError:
            catalog = Catalog.empty()
    lang: Lang = body["lang"]
    meta = meta or DocumentMeta()
    resolved = resolve(body, values=values or {}, blocks=blocks or {}, catalog=catalog)
    page: Mapping[str, Any] = body.get("page") or {}
    tuning_raw: Mapping[str, Any] = page.get("tuning") or {}
    setup = page_setup(
        layout,
        Orientation(page.get("orientation", "portrait")),
        Density(page["density"]) if page.get("density") else None,
        Tuning(
            line_height=tuning_raw.get("lineHeight"),
            margin_delta_mm=int(tuning_raw.get("marginDeltaMm", 0)),
            lead_mm=int(tuning_raw.get("leadMm", 0)),
        ),
    )
    warnings = resolved.warnings()
    document = Document()
    _page(document, setup)
    _base_style(document, letterhead, setup)
    writer = _DocxWriter(document, lang, letterhead, setup)
    warnings += _headers(document, writer, letterhead, meta, layout, page_numbers)
    _letter_meta(document, writer, meta, layout)
    if meta.title:
        title = document.add_paragraph()
        run = title.add_run(meta.title)
        run.bold = True
        run.font.size = Pt(round(setup.font_pt * 1.4, 1))
        run.font.color.rgb = writer.accent
    writer.blocks(resolved.content)
    if meta.executor and layout is not Layout.COMMERCIAL:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(24)
        run = paragraph.add_run(EXECUTOR[lang] + meta.executor)
        run.font.size = Pt(7.5)
        run.font.color.rgb = _MUTED
    out = BytesIO()
    document.save(out)
    return DocxResult(out.getvalue(), tuple(warnings))


def _page(document: DocxDocument, setup: PageSetup) -> None:
    section = document.sections[0]
    top, right, bottom, left = setup.margins_mm
    if setup.orientation is Orientation.LANDSCAPE:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(297), Mm(210)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.right_margin = Mm(top + setup.lead_mm), Mm(right)
    section.bottom_margin, section.left_margin = Mm(bottom), Mm(left)


def _base_style(document: DocxDocument, letterhead: Letterhead, setup: PageSetup) -> None:
    style: Any = document.styles["Normal"]
    style.font.name = _FONT[letterhead.font]
    style.font.size = Pt(setup.font_pt)
    fonts = style.element.rPr.rFonts
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), _FONT[letterhead.font])
    style.paragraph_format.line_spacing = setup.line_height
    style.paragraph_format.space_after = Pt(round(setup.font_pt * 0.55, 1))
    style.paragraph_format.widow_control = True


def _headers(
    document: DocxDocument,
    writer: _DocxWriter,
    letterhead: Letterhead,
    meta: DocumentMeta,
    layout: Layout,
    page_numbers: bool,
) -> list[str]:
    warnings: list[str] = []
    section = document.sections[0]
    section.different_first_page_header_footer = True
    first = section.first_page_header.paragraphs[0]
    if layout is not Layout.INTERNAL and letterhead.logo is not None:
        if letterhead.logo.mime in ("image/png", "image/jpeg"):
            first.add_run().add_picture(BytesIO(letterhead.logo.data), height=Mm(12))
            first = section.first_page_header.add_paragraph()
        else:
            warnings.append("docx: logo format is not supported by python-docx (png or jpeg needed)")
    name = first.add_run(letterhead.company_name)
    name.bold = True
    name.font.size = Pt(13 if layout is not Layout.INTERNAL else 10)
    name.font.color.rgb = writer.accent
    if layout is not Layout.INTERNAL:
        for line in letterhead.requisites:
            run = section.first_page_header.add_paragraph().add_run(line)
            run.font.size = Pt(7.5)
            run.font.color.rgb = _MUTED
    if meta.number:
        other = section.header.paragraphs[0]
        other.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = other.add_run(meta.number)
        run.font.size = Pt(7.5)
        run.font.color.rgb = _MUTED
    for footer in (section.first_page_footer, section.footer):
        paragraph = footer.paragraphs[0]
        if layout is Layout.COMMERCIAL:
            for index, line in enumerate(letterhead.footer):
                target = paragraph if index == 0 else footer.add_paragraph()
                run = target.add_run(line)
                run.font.size = Pt(6.5)
                run.font.color.rgb = _MUTED
        if page_numbers:
            numbers = footer.add_paragraph() if layout is Layout.COMMERCIAL and letterhead.footer else paragraph
            numbers.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            before, middle, after = PAGE_OF[writer.lang]
            numbers.add_run(before)
            _field_run(numbers, "PAGE")
            numbers.add_run(middle)
            _field_run(numbers, "NUMPAGES")
            if after:
                numbers.add_run(after)
            for run in numbers.runs:
                run.font.size = Pt(7.5)
    return warnings


def _letter_meta(document: DocxDocument, writer: _DocxWriter, meta: DocumentMeta, layout: Layout) -> None:
    if layout is Layout.COMMERCIAL:
        return
    lang = writer.lang
    left = number_line(lang, meta.number, format_date(meta.date, lang, long=False) if meta.date else None)
    if not (left or meta.reply_to or meta.addressee):
        return
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.text = left
    if meta.reply_to:
        cell.add_paragraph(REPLY_TO[lang] + meta.reply_to)
    right = table.cell(0, 1)
    right.text = "\n".join(meta.addressee)
    document.add_paragraph()
