"""Экспорт DOCX из того же тела: текст, таблицы с повтором шапки, ориентация, колонтитулы бланка."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from rt_docs import DocumentMeta, Layout, to_docx
from tests import samples


def open_docx(data: bytes) -> Any:
    return Document(BytesIO(data))


def all_text(doc: Any) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts += [cell.text for cell in row.cells]
    return "\n".join(parts)


def test_letter_text_values_signature_and_executor() -> None:
    result = to_docx(
        samples.letter_short(),
        letterhead=samples.LETTERHEAD,
        layout=Layout.LETTER,
        meta=samples.LETTER_META,
        values=samples.VALUES,
        page_numbers=False,
    )
    doc = open_docx(result.data)
    text = all_text(doc)
    assert "Уважаемый Бахтиёр Бахтиярович!" in text
    assert "[Номер заявки]" in text  # неразрешённая переменная — подписью из каталога
    assert "А. А. Демонстрационный" in text and "М. П." in text
    assert "Исп.: Исполнитель А. Б. Петрова" in text
    assert "№ ИСХ-015/26 от 28.10.2026" in text
    header = "\n".join(p.text for p in doc.sections[0].first_page_header.paragraphs)
    assert "ООО «Демо Рельс Сервис»" in header
    assert result.warnings == ("variable without value: deal.number",)
    footer_xml = doc.sections[0].footer._element.xml
    assert "NUMPAGES" not in footer_xml


def test_invoice_table_repeats_header_and_rows_do_not_split() -> None:
    blocks = {
        "doc.items_table": samples.items_table(40),
        "doc.totals": samples.totals("Сто пятьдесят два миллиона триста сорок тысяч сум 00 тийин"),
        "doc.payment_requisites": samples.payment_requisites(),
    }
    result = to_docx(
        samples.invoice(),
        letterhead=samples.LETTERHEAD,
        layout=Layout.COMMERCIAL,
        meta=DocumentMeta(title="Счёт на оплату № 47", number="СЧ-47/26"),
        values=samples.VALUES,
        blocks=blocks,
    )
    doc = open_docx(result.data)
    items = doc.tables[0]
    assert len(items.rows) == 41
    first_row = items.rows[0]._tr
    assert first_row.trPr.find(qn("w:tblHeader")) is not None
    assert items.rows[5]._tr.trPr.find(qn("w:cantSplit")) is not None
    assert items.rows[5]._tr.trPr.find(qn("w:tblHeader")) is None
    assert "Позиция 40" in all_text(doc) and "Итого к оплате" in all_text(doc)
    assert "NUMPAGES" in doc.sections[0].footer._element.xml
    assert "СЧ-47/26" in "\n".join(p.text for p in doc.sections[0].header.paragraphs)


def test_specification_is_landscape() -> None:
    result = to_docx(
        samples.specification(),
        letterhead=samples.LETTERHEAD,
        layout=Layout.COMMERCIAL,
        values=samples.VALUES,
        blocks={"doc.spec_table": samples.items_table(3, landscape=True)},
    )
    section = open_docx(result.data).sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width > section.page_height


def test_lists_links_headings_and_page_break() -> None:
    body = samples.body(
        [
            samples.heading(1, "Заголовок"),
            {
                "type": "bulletList",
                "content": [{"type": "listItem", "content": [samples.para("Пункт")]}],
            },
            {
                "type": "orderedList",
                "attrs": {"start": 3},
                "content": [{"type": "listItem", "content": [samples.para("Три")]}],
            },
            samples.para(
                {"type": "text", "text": "сайт", "marks": [{"type": "link", "attrs": {"href": "https://example.org"}}]}
            ),
            {"type": "pageBreak"},
            samples.para(samples.text("жирный", "bold", "italic")),
        ],
        lang="uz",
    )
    doc = open_docx(to_docx(body, letterhead=samples.LETTERHEAD, layout=Layout.INTERNAL).data)
    xml = doc.element.xml
    assert "Заголовок" in xml and 'w:type="page"' in xml and "w:hyperlink" in xml
    styles = [p.style.name for p in doc.paragraphs]
    assert "List Bullet" in styles and "List Number" in styles
    bold = next(r for p in doc.paragraphs for r in p.runs if r.text == "жирный")
    assert bold.bold and bold.italic
