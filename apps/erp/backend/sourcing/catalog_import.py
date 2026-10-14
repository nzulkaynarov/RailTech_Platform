"""Self-service seller catalog import engine (true-PIM).

A seller uploads an .xlsx of their products + commercial terms. We:
  * resolve/create the **Brand** (catalog.Manufacturer) by name,
  * de-duplicate the canonical **Product** by (brand + manufacturer part number)
    or GOST/drawing — so two suppliers of the same item share ONE PIM card,
  * upsert the seller's **SupplierOffer** (price / currency / stock / lead / MOQ).

Файл несёт и ТАКСОНОМИЮ карточки — категорию, код ТН ВЭД и направление. Без них загрузка
физически не могла дать публикуемую карточку: порог публикации (`catalog/completeness.py`)
требует ровно эти поля, и прайс-лист на 3000 строк давал 3000 черновиков, о чём отчёт молчал.
Поэтому же отчёт теперь считает, сколько карточек порог не пройдут и чего им не хватает.

The supplier is the UPLOADER's own company (derived from membership in the
view) — never a column in the file. Each row is processed in its own savepoint
so one bad row can't poison the batch; per-row problems are returned as a
structured error list. `dry_run=True` runs the full pipeline and rolls back,
returning the same counts for a preview.

Business logic lives here (Rule 8); the Celery task is a thin wrapper.
"""

import logging
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO

import pandas as pd
from django.db import transaction

logger = logging.getLogger(__name__)

MAX_EXCEL_ROWS = 50_000
MAX_REPORTED_ERRORS = 200
#: Предупреждения живут по тем же правилам, что и ошибки: показываем первые N, считаем все.
MAX_REPORTED_WARNINGS = MAX_REPORTED_ERRORS

# Logical field -> substring aliases (matched against lower-cased headers).
# Order matters: more specific fields claim their column before generic ones
# (e.g. "артикул производителя" is taken by `mpn`, leaving bare "артикул" for
# `supplier_sku`).
# NOTE: resolution order matters — the more specific identifier columns (mpn,
# gost, supplier_sku) claim their column BEFORE `brand`, because the brand alias
# "производитель" is a substring of "артикул производителя" (MPN) and would
# otherwise steal it.
COLUMN_ALIASES = {
    # `category` стоит ПЕРВЫМ намеренно. Алиас «категор» не встречается ни в одном заголовке
    # остальных полей, так что ничего у них он не отнимает; зато заголовок «Категория товара»
    # иначе достался бы полю `name` (у него алиас «товар»), и карточки назвались бы категориями.
    "category": ["категор", "category"],
    "name": ["name", "наименование", "товар", "название", "изделие"],
    "mpn": ["артикул производителя", "mpn", "part number", "part_no", "заводской", "номер детали"],
    "gost": ["гост", "gost", "чертеж", "drawing"],
    "supplier_sku": ["артикул поставщика", "supplier sku", "ваш артикул", "артикул продавца", "артикул"],
    "brand": ["бренд", "brand", "производитель", "manufacturer", "марка"],
    "description": ["описание", "desc", "примечание"],
    "uom": ["ед", "изм", "uom", "unit"],
    "price": ["цена", "price", "стоимость"],
    "currency": ["валюта", "currency"],
    "stock": ["наличие", "склад", "stock"],
    "lead": ["срок", "lead", "поставк"],
    "moq": ["мин", "moq", "минимальн"],
    # Остальные новинки — в САМОМ КОНЦЕ: так они могут занять только колонку, которую не
    # забрал никто до них, и разбор существующих файлов не меняется. Алиасы узкие осознанно:
    # общее «код» отобрало бы у артикула «Код поставщика», а «направление» безопасно только
    # целиком — с «наименованием» оно не пересекается ни одной подстрокой.
    "hs_code": ["тн вэд", "тнвэд", "тн-вэд", "hs code", "hs_code", "hs-code", "hscode"],
    "domain_type": ["направление", "домен", "domain"],
}


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    offers_upserted: int = 0
    skipped: int = 0
    errors: list = field(default_factory=list)
    #: Мягкие замечания: строка ЗАГРУЖЕНА, но к ней есть вопрос (код ТН ВЭД вне справочника).
    #: Предупреждение не откатывает строку и не попадает в `skipped` — иначе это был бы блок.
    warnings: list = field(default_factory=list)
    warning_count: int = 0
    #: Сколько строк файла не попало в разбор из-за `MAX_EXCEL_ROWS` (раньше исчезали молча).
    rows_truncated: int = 0
    #: Сколько затронутых карточек НЕ пройдут порог публикации и чего им не хватает.
    unpublishable: int = 0
    missing_breakdown: list = field(default_factory=list)
    dry_run: bool = False

    def add_error(self, row_number: int, message: str) -> None:
        self.skipped += 1
        if len(self.errors) < MAX_REPORTED_ERRORS:
            self.errors.append({"row": row_number, "message": message})

    def add_warning(self, row_number: int, message: str) -> None:
        self.warning_count += 1
        if len(self.warnings) < MAX_REPORTED_WARNINGS:
            self.warnings.append({"row": row_number, "message": message})

    def as_dict(self) -> dict:
        return {
            "status": "success",
            "dry_run": self.dry_run,
            "created": self.created,
            "updated": self.updated,
            "offers_upserted": self.offers_upserted,
            "skipped": self.skipped,
            # `error_count`/`warning_count` считают ВСЕ находки, а списки усечены до
            # MAX_REPORTED_*: интерфейсу этого достаточно, чтобы сказать «показаны первые 200
            # из N», а не делать вид, что их всего 200.
            "error_count": self.skipped,
            "errors": self.errors,
            "warning_count": self.warning_count,
            "warnings": self.warnings,
            "rows_truncated": self.rows_truncated,
            "unpublishable": self.unpublishable,
            "missing_breakdown": self.missing_breakdown,
            # Backward-compatible key consumed by the existing frontend.
            "items_processed": self.created + self.updated,
        }


def _resolve_columns(columns) -> dict:
    """Map each logical field to an actual column, claiming each column once."""
    resolved = {}
    claimed = set()
    for logical, aliases in COLUMN_ALIASES.items():
        for column in columns:
            if column in claimed:
                continue
            if any(alias in column for alias in aliases):
                resolved[logical] = column
                claimed.add(column)
                break
        else:
            resolved[logical] = None
    return resolved


def _cell(row, column) -> str:
    """Trimmed string for a cell; '' for missing/NaN."""
    if not column:
        return ""
    value = row.get(column)
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _parse_decimal(text: str):
    """Parse a price cell. Returns None for a blank cell (name-only row → no offer).

    A non-blank value must be a number > 0 — a negative or zero price is a data
    error, not a valid offer (M3). We reject it with a stable code so the row is
    rolled back and reported instead of silently writing a bogus offer.
    """
    if not text:
        return None
    try:
        value = Decimal(text.replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ValueError("PRICE_NOT_A_NUMBER")
    if value <= 0:
        raise ValueError("PRICE_NOT_POSITIVE")
    return value


def _parse_int(text: str, default: int) -> int:
    if not text:
        return default
    try:
        return int(float(text.replace(" ", "").replace(",", ".")))
    except (ValueError, TypeError):
        return default


def _hs_digits(text: str) -> str:
    """Цифры кода ТН ВЭД из ячейки.

    Excel отдаёт числовую ячейку как `8482100000.0` — дробный хвост принадлежит представлению
    числа, а не коду, и без его отсечения десятизначный код превращался бы в одиннадцатизначный.
    """
    text = (text or "").strip()
    whole, dot, frac = text.partition(".")
    if dot and whole.isdigit() and set(frac) <= {"0"}:
        text = whole
    return "".join(char for char in text if char.isdigit())


@dataclass
class _Ctx:
    """Per-import shared state (caches + lookups) — avoids per-row DB queries."""

    supplier: object
    Product: object
    Manufacturer: object
    SupplierOffer: object
    uom: dict
    currency: dict
    stock: dict
    domain: dict
    #: Домен по КОДУ и по НАЗВАНИЮ в нижнем регистре — продавец пишет в файле любое из двух.
    domain_by_key: dict
    #: Категория по названию в нижнем регистре. Таксономия платформенная: импорт её ЧИТАЕТ.
    category_by_name: dict
    #: Вердикт справочника ТН ВЭД по каждому коду файла — считается пачкой ДО обхода строк.
    hs_status: dict
    hs_code_max_length: int
    status_draft: object
    brand_cache: dict
    #: Затронутые карточки по pk (одна и та же строка может встретиться дважды) — по ним
    #: считается сводка публикуемости.
    touched_products: dict
    dry_run: bool


def _resolve_brand(ctx: _Ctx, name: str):
    """Find a Brand by name, creating it on demand (catalog grows from uploads).

    In dry-run we never create — an unseen brand resolves to None, which only
    softens product de-dup for the preview.
    """
    if not name:
        return None
    key = name.lower()
    if key in ctx.brand_cache:
        return ctx.brand_cache[key]
    brand = ctx.Manufacturer.objects.filter(name__iexact=name).first()
    if not brand and not ctx.dry_run:
        # Тот же нормализатор, что у формы карточки (catalog.services.find_or_create_manufacturer).
        from catalog.services import find_or_create_manufacturer

        brand = find_or_create_manufacturer(name)
    ctx.brand_cache[key] = brand
    return brand


def _find_existing_product(ctx: _Ctx, brand, mpn: str, gost: str):
    """De-dup key: (brand + manufacturer part number), else GOST/drawing.

    `select_related("domain_type")` — не украшение: домен читает и порог публикации, и
    сводка отчёта, и без него каждая обновлённая карточка стоила бы лишний запрос.
    """
    query = ctx.Product.objects.select_related("domain_type")
    if brand and mpn:
        return query.filter(brand=brand, manufacturer_part_number__iexact=mpn).first()
    if gost:
        query = query.filter(gost_drawing__iexact=gost)
        if brand:
            query = query.filter(brand=brand)
        return query.first()
    return None


def _save_product(product) -> None:
    product.save()


def _update_product(ctx: _Ctx, existing, fields: dict, brand, uom) -> None:
    """Обновить существующую карточку значениями строки (пустая ячейка ничего не стирает)."""
    existing.name = fields["name"]
    if fields["description"]:
        existing.description = fields["description"]
    if brand:
        existing.brand = brand
    if fields["mpn"]:
        existing.manufacturer_part_number = fields["mpn"]
    if fields["gost"]:
        existing.gost_drawing = fields["gost"]
    if uom:
        existing.uom = uom
    if fields["category"]:
        existing.category = fields["category"]
    if fields["hs_code"]:
        existing.hs_code = fields["hs_code"]
    # ЯВНО указанное направление применяется и при обновлении: иначе домен карточки, однажды
    # угаданный по ключевым словам, нельзя было исправить ничем, кроме ручной правки. ДОГАДКА
    # же остаётся только для создания — молча переклассифицировать чужую готовую карточку по
    # совпадению слова в наименовании нельзя.
    if fields["domain"]:
        existing.domain_type = fields["domain"]
    # Backfill ownership for an orphaned card (C3), but never hijack a product
    # another tenant already owns — de-dup matches globally on brand+MPN.
    if ctx.supplier and existing.owner_company_id is None:
        existing.owner_company = ctx.supplier
    _save_product(existing)


def _create_product(ctx: _Ctx, fields: dict, brand, uom):
    from .tasks import route_domain

    domain = fields["domain"] or ctx.domain.get(route_domain(fields["name"]))
    product = ctx.Product(
        name=fields["name"],
        description=fields["description"],
        brand=brand,
        manufacturer_part_number=fields["mpn"],
        gost_drawing=fields["gost"],
        domain_type=domain,
        category=fields["category"],
        hs_code=fields["hs_code"],
        status=ctx.status_draft,
        uom=uom,
        # Stamp the importing seller as owner so they can see/manage their own cards
        # (C3) — without this, imported products are orphaned and invisible to them.
        owner_company=ctx.supplier,
    )
    _save_product(product)
    return product


def _apply_product(ctx: _Ctx, fields: dict):
    """Create or update the canonical Product card. Returns (product, created)."""
    brand = _resolve_brand(ctx, fields["brand_name"])
    existing = _find_existing_product(ctx, brand, fields["mpn"], fields["gost"])
    uom = ctx.uom.get(fields["uom_code"]) if fields["uom_code"] else None

    if existing:
        _update_product(ctx, existing, fields, brand, uom)
        return existing, False
    return _create_product(ctx, fields, brand, uom), True


def _upsert_offer(ctx: _Ctx, product, offer: dict) -> None:
    """Upsert the uploader's SupplierOffer for this product (unique per supplier)."""
    ctx.SupplierOffer.objects.update_or_create(
        product=product,
        supplier=ctx.supplier,
        defaults={
            "base_price": offer["price"],
            "currency": offer["currency"],
            "stock_status": offer["stock"],
            "supplier_sku": offer["supplier_sku"],
            "delivery_lead_time": offer["lead"],
            "min_order_qty": offer["moq"],
            "is_active": True,
        },
    )


def _row_taxonomy(ctx: _Ctx, row, columns: dict):
    """Категория / код ТН ВЭД / направление строки. Возвращает `(значения, предупреждения)`.

    Категория и направление СТРОГИЕ: продавец написал слово, которого в справочнике нет —
    молча проглотить его значит соврать. Категорию при этом импорт только ЧИТАЕТ: таксономия
    платформенная, продавец в неё не пишет (это продуктовое решение, а не недосмотр).

    Код ТН ВЭД, наоборот, МЯГКИЙ (PRODUCT RULE 2): классификатор может быть ещё не загружен, а
    код продавец знает из своей декларации — расхождение со справочником это ПРЕДУПРЕЖДЕНИЕ.
    """
    category_name = _cell(row, columns["category"])
    category = ctx.category_by_name.get(category_name.lower()) if category_name else None
    if category_name and category is None:
        raise ValueError(f"UNKNOWN_CATEGORY:{category_name}")

    domain_text = _cell(row, columns["domain_type"])
    domain = ctx.domain_by_key.get(domain_text.lower()) if domain_text else None
    if domain_text and domain is None:
        raise ValueError(f"UNKNOWN_DOMAIN:{domain_text}")

    warnings = []
    hs_code = _hs_digits(_cell(row, columns["hs_code"]))
    if hs_code and len(hs_code) > ctx.hs_code_max_length:
        # Длиннее поля карточки — записать нельзя, но и ронять строку из-за этого не за что.
        warnings.append(f"HS_CODE_TOO_LONG:{hs_code}")
        hs_code = ""
    elif hs_code:
        verdict = ctx.hs_status.get(hs_code)
        if verdict and not verdict["known"]:
            warnings.append(f"HS_CODE_NOT_IN_CLASSIFIER:{hs_code}")
        elif verdict and not verdict["active"]:
            warnings.append(f"HS_CODE_RETIRED:{hs_code}")

    return {"category": category, "domain": domain, "hs_code": hs_code}, warnings


def _row_commercials(ctx: _Ctx, row, columns: dict):
    """Валюта / наличие / цена строки. Коды сверяются со справочниками строго."""
    currency_code = _cell(row, columns["currency"]).upper()
    currency = ctx.currency.get(currency_code) if currency_code else None
    if currency_code and not currency:
        raise ValueError(f"UNKNOWN_CURRENCY:{currency_code}")

    stock_code = _cell(row, columns["stock"]).upper()
    stock = ctx.stock.get(stock_code) if stock_code else None
    if stock_code and not stock:
        raise ValueError(f"UNKNOWN_STOCK_STATUS:{stock_code}")

    return currency, stock, _parse_decimal(_cell(row, columns["price"]))


def _process_one_row(ctx: _Ctx, row, columns: dict):
    """Handle a single row. Returns an outcome dict, or None for a blank row.

    Raises ValueError with a stable code for row-level problems so the caller
    can roll back this row's savepoint and report it.
    """
    name = _cell(row, columns["name"])
    if not name:
        return None

    currency, stock, price = _row_commercials(ctx, row, columns)
    taxonomy, warnings = _row_taxonomy(ctx, row, columns)

    product, created = _apply_product(
        ctx,
        {
            "name": name,
            "description": _cell(row, columns["description"]),
            "brand_name": _cell(row, columns["brand"]),
            "mpn": _cell(row, columns["mpn"]),
            "gost": _cell(row, columns["gost"]),
            "uom_code": _cell(row, columns["uom"]).lower(),
            **taxonomy,
        },
    )

    offer_written = False
    if ctx.supplier and price is not None:
        _upsert_offer(
            ctx,
            product,
            {
                "price": price,
                "currency": currency,
                "stock": stock,
                "supplier_sku": _cell(row, columns["supplier_sku"]),
                "lead": _cell(row, columns["lead"]),
                "moq": _parse_int(_cell(row, columns["moq"]), 1),
            },
        )
        offer_written = True

    return {"created": created, "offer": offer_written, "product": product, "warnings": warnings}


def _domain_keys(domains) -> dict:
    """Направление и по коду (`GENERAL_GOODS`), и по человеческому названию — регистр не важен.

    Название кладём первым, код — вторым: при совпадении выигрывает КОД, он однозначен.
    """
    by_key = {domain.name.lower(): domain for domain in domains if domain.name}
    by_key.update({domain.code.lower(): domain for domain in domains})
    return by_key


def _build_context(supplier, dry_run: bool, hs_status: dict = None) -> _Ctx:
    from catalog.models import Category, DomainType, Manufacturer, Product, ProductStatus, StockStatus, SupplierOffer
    from master_data.models import Currency, UnitOfMeasure

    domains = list(DomainType.objects.all())
    return _Ctx(
        supplier=supplier,
        Product=Product,
        Manufacturer=Manufacturer,
        SupplierOffer=SupplierOffer,
        uom={u.code.lower(): u for u in UnitOfMeasure.objects.all()},
        currency={c.code.upper(): c for c in Currency.objects.all()},
        stock={s.code.upper(): s for s in StockStatus.objects.all()},
        domain={d.code: d for d in domains},
        domain_by_key=_domain_keys(domains),
        # Только ДЕЙСТВУЮЩИЕ категории: выключенная ветка таксономии — не цель для новых карточек.
        # `order_by` фиксирует победителя при одинаковых названиях в разных ветках дерева.
        category_by_name={
            category.name.strip().lower(): category
            for category in Category.objects.filter(is_active=True).order_by("-sort_order", "-id")
        },
        hs_status=hs_status or {},
        hs_code_max_length=Product._meta.get_field("hs_code").max_length,
        status_draft=ProductStatus.objects.filter(code="DRAFT").first(),
        brand_cache={},
        touched_products={},
        dry_run=dry_run,
    )


def _iterate_rows(ctx: _Ctx, data_frame, columns: dict, result: ImportResult, progress_callback) -> None:
    total = len(data_frame)
    for index, row in data_frame.iterrows():
        row_number = index + 2  # +1 for header, +1 for 1-based display
        try:
            with transaction.atomic():
                outcome = _process_one_row(ctx, row, columns)
        except ValueError as exc:
            result.add_error(row_number, str(exc))
            continue
        except Exception:  # noqa: BLE001 — isolate one row, keep importing
            logger.exception("Catalog import row %s failed", row_number)
            result.add_error(row_number, "ROW_FAILED")
            continue

        if outcome is None:
            continue
        if outcome["created"]:
            result.created += 1
        else:
            result.updated += 1
        if outcome["offer"]:
            result.offers_upserted += 1
        # Предупреждения записываются ТОЛЬКО после успеха строки: у откатанной строки их нет.
        for message in outcome["warnings"]:
            result.add_warning(row_number, message)
        ctx.touched_products[outcome["product"].pk] = outcome["product"]

        processed = result.created + result.updated + result.skipped
        if progress_callback and processed % 500 == 0:
            progress_callback(processed, total)


def _count_file_rows(file_content: bytes) -> int:
    """Сколько СОДЕРЖАТЕЛЬНЫХ строк данных в первом листе (без строки заголовка).

    Считаем потоком (`read_only=True` у openpyxl отдаёт строки по одной и не поднимает лист
    в память целиком) и только когда файл заведомо длиннее предела: платить за второй проход
    ради числа «сколько отрезано» имеет смысл лишь тогда, когда отрезать действительно есть что.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(file_content), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        filled = sum(
            1
            for cells in sheet.iter_rows(values_only=True)
            if any(cell is not None and str(cell).strip() for cell in cells)
        )
    finally:
        workbook.close()
    return max(0, filled - 1)  # первая содержательная строка — заголовок


def _read_rows(file_content: bytes):
    """`(DataFrame, сколько строк файла не попало в разбор)`.

    `nrows=MAX_EXCEL_ROWS` раньше отрезал хвост МОЛЧА: продавец видел «создано 50 000» и не
    узнавал, что остальные строки его прайса система просто не читала. Читаем на одну строку
    больше предела — лишняя строка и есть признак переполнения.

    `dtype=str` — НЕ косметика. Без него pandas выводит тип колонки сам, и колонка кодов ТН ВЭД
    из текстовой становится числовой: `0101210000` превращается в `101210000`, теряя ведущий
    ноль. Продавец получал замечание «кода нет в справочнике» на совершенно верный код, а в
    карточку писался испорченный — при том что нули начинают целые группы классификатора
    (01–09). Ловится только глазами: файл разбирается «успешно», просто цифра пропадает.
    Ячейки и так разбираются нашими парсерами из строк, поэтому чтение всего как текста ничего
    больше не меняет.
    """
    data_frame = pd.read_excel(BytesIO(file_content), nrows=MAX_EXCEL_ROWS + 1, dtype=str)
    if len(data_frame) <= MAX_EXCEL_ROWS:
        return data_frame, 0
    return data_frame.head(MAX_EXCEL_ROWS), max(0, _count_file_rows(file_content) - MAX_EXCEL_ROWS)


def _prescan_hs_codes(data_frame, columns: dict) -> dict:
    """Вердикт классификатора по ВСЕМ кодам файла — два запроса на импорт вместо двух на строку."""
    column = columns.get("hs_code")
    if not column:
        return {}

    from catalog.hs_suggest import validate_hs_codes

    values = data_frame[column]
    if getattr(values, "ndim", 1) > 1:  # два заголовка совпали после нормализации регистра
        values = values.iloc[:, 0]
    codes = {_hs_digits(str(value)) for value in values.tolist()}
    return validate_hs_codes(codes)


def _summarize_publishability(ctx: _Ctx, result: ImportResult) -> None:
    """Дописать в отчёт, сколько карточек НЕ пройдут порог публикации и чего им не хватает."""
    from catalog.completeness import summarize_completeness

    summary = summarize_completeness(ctx.touched_products.values())
    result.unpublishable = summary["unpublishable"]
    result.missing_breakdown = summary["missing_breakdown"]


def run_catalog_import_frame(data_frame, supplier_company_id=None, dry_run=False, progress_callback=None) -> dict:
    """Ядро импорта: DataFrame строк → upsert PIM-карточек + офферов загрузившего.

    Выделено из байтовой версии (R5b), чтобы ИИ-путь (`services_ai_catalog` собирает строки из
    PDF/фото/произвольного файла) проходил через ТОТ ЖЕ разбор шапки, те же справочники и тот же
    отчёт, а не через параллельный движок. Заголовки колонок нормализуются здесь же, поэтому на
    вход годится и DataFrame прямо из файла, и DataFrame из словарей с русскими заголовками
    шаблона. Returns a structured result dict (see ImportResult.as_dict). Raises
    ValueError('NAME_COLUMN_NOT_FOUND') when the frame has no recognizable name
    column — the only fatal, file-level error.
    """
    from accounts.models import Company

    supplier = Company.objects.filter(id=supplier_company_id).first() if supplier_company_id else None

    data_frame.columns = [str(column).lower().strip() for column in data_frame.columns]
    columns = _resolve_columns(list(data_frame.columns))
    if not columns["name"]:
        raise ValueError("NAME_COLUMN_NOT_FOUND")

    ctx = _build_context(supplier, dry_run, hs_status=_prescan_hs_codes(data_frame, columns))
    result = ImportResult(dry_run=dry_run)

    if dry_run:
        with transaction.atomic():
            _iterate_rows(ctx, data_frame, columns, result, progress_callback)
            # Сводка считается ВНУТРИ транзакции: карточки предпросмотра живут только в ней.
            _summarize_publishability(ctx, result)
            transaction.set_rollback(True)
    else:
        _iterate_rows(ctx, data_frame, columns, result, progress_callback)
        _summarize_publishability(ctx, result)

    return result.as_dict()


def run_catalog_import(file_content: bytes, supplier_company_id=None, dry_run=False, progress_callback=None) -> dict:
    """Parse an .xlsx and upsert PIM products + the uploader's offers.

    Тонкая байтовая обёртка над `run_catalog_import_frame`: читает файл (с пределом строк) и
    дописывает в отчёт `rows_truncated` — единственное, что знает только файловый слой.
    """
    data_frame, rows_truncated = _read_rows(file_content)
    result = run_catalog_import_frame(
        data_frame,
        supplier_company_id=supplier_company_id,
        dry_run=dry_run,
        progress_callback=progress_callback,
    )
    result["rows_truncated"] = rows_truncated
    return result
